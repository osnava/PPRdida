"""Gráficas y tablas del ejercicio PPR vs VUAA (leyendo sim_ppr_vs_vuaa)."""
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.ticker as mtick
from matplotlib.lines import Line2D
from matplotlib.patches import Patch
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from sim_ppr_vs_vuaa import APORT_M, ESCENARIOS, MESES, REFUND

BASE = Path(__file__).resolve().parent.parent   # raíz del proyecto
OUT = BASE / "outputs" / "graficas"
OUT.mkdir(parents=True, exist_ok=True)
M = 1e6
AZUL, NARANJA, VERDE, VERDE_CLARO, ROJO, GRIS = "#1f77b4", "#ff7f0e", "#2ca02c", "#74c476", "#d62728", "#7f7f7f"
plt.rcParams.update({"font.size": 10, "axes.spines.top": False, "axes.spines.right": False})

grupos = [  # (título, [índices en ESCENARIOS])
    ("Peso estable (S&P 10% anual en MXN)", [0, 1, 2, 3]),
    ("Peso depreciándose 1.5%/año (S&P 11.65% en MXN)", [4, 5, 6, 7]),
]
ESTILOS = [  # por posición dentro del grupo: (color, etiqueta corta, estilo de línea)
    (AZUL, "PPR · cargos al tope", "-"), (NARANJA, "PPR · mitad de tope", "-"),
    (VERDE, "VUAA · mismos aportes que el PPR", "-"), (VERDE_CLARO, "VUAA · mismo bolsillo que el PPR", "--"),
]
estilo = {i: ESTILOS[i % 4] for i in range(len(ESCENARIOS))}
orden = [2, 3, 1, 0, 6, 7, 5, 4]  # por grupo, de mayor a menor neto


def etiqueta_barra(i):
    return f"{estilo[i][1]}\n{grupos[i // 4][0].split(' (')[0]}"


# ---------------------------------------------------------------- 1. evolución
fig, axes = plt.subplots(1, 2, figsize=(13, 5.6), sharey=True)
anios = np.arange(1, MESES // 12 + 1)
for ax, (titulo, idxs) in zip(axes, grupos):
    etiquetas = []
    for idx in idxs:
        nombre, r = ESCENARIOS[idx]
        color, etiqueta, ls = estilo[idx]
        serie = np.array(r["serie"]) / M
        ax.plot(anios, serie, color=color, lw=2.2, ls=ls, label=etiqueta)
        neto_fin = r["neto"] / M
        ax.plot([40, 40], [neto_fin, serie[-1]], color=ROJO, lw=2.0, ls=":")  # venta y/o ISR
        ax.plot([40], [neto_fin], "o", color=color, ms=6)
        etiquetas.append([neto_fin, color, neto_fin])   # [y etiqueta, color, valor]
    # anti-solape: separación mínima de 3.5M entre etiquetas y fondo blanco;
    # cada etiqueta conserva su propio valor y color
    etiquetas.sort(key=lambda t: -t[0])
    for j in range(1, len(etiquetas)):
        if etiquetas[j][0] > etiquetas[j - 1][0] - 3.5:
            etiquetas[j][0] = etiquetas[j - 1][0] - 3.5
    for yv, color, val in etiquetas:
        ax.text(40.4, yv, f"neto \\${val:.1f}M", fontsize=9, color=color, fontweight="bold", va="center",
                bbox=dict(facecolor="white", alpha=0.75, edgecolor="none", pad=1))
    bono = ESCENARIOS[idxs[0]][1]["bono_acreditado"]
    ax.axvline(25, ls="--", color=GRIS, lw=1)
    ax.text(25.4, ax.get_ylim()[1] * 0.05, f"año 25: se acredita\nel Bono (\\${bono / 1e3:.0f}k)", fontsize=8, color=GRIS)
    ax.set_title(titulo, fontsize=10.5)
    ax.set_xlabel("Años")
    ax.yaxis.set_major_formatter(mtick.StrMethodFormatter("{x:,.0f} M"))
    ax.set_xlim(1, 46)
    ax.tick_params(labelleft=True)
    handles, labels = ax.get_legend_handles_labels()
    handles += [Line2D([], [], color=ROJO, lw=2, ls=":"),
                Line2D([], [], color=GRIS, lw=1, ls="--")]
    labels += ["Caída al retiro: corretaje (VUAA) e ISR", "Año 25: bono PPR acreditado"]
    ax.legend(handles, labels, loc="upper left", frameon=False, fontsize=8.2)
fig.suptitle("Saldo antes de retiro (línea) y neto tras venta e ISR (punto) — PPR vs VUAA · 5,000 MXN/mes · 40 años",
             fontsize=11, fontweight="bold")
axes[0].set_ylabel("Millones MXN (nominales)")
fig.tight_layout()
fig.savefig(OUT / "01_evolucion_saldos.png", dpi=150)
plt.close(fig)

# ---------------------------------------------------- 2. neto tras impuesto
fig, ax = plt.subplots(figsize=(11.5, 7))
y = np.arange(len(orden))
netos = [ESCENARIOS[i][1]["neto"] / M for i in orden]
imps = [ESCENARIOS[i][1]["impuesto"] / M for i in orden]
colores = [estilo[i][0] for i in orden]
ax.barh(y, netos, color=colores)
ax.barh(y, imps, left=netos, color=ROJO, alpha=0.85)
for yi, n, im in zip(y, netos, imps):
    ax.text(n + im + 0.6, yi, f"neto \\${n:.1f}M  (ISR \\${im:.1f}M)", va="center", fontsize=9)
ax.set_yticks(y)
ax.set_yticklabels([etiqueta_barra(i) for i in orden], fontsize=9)
ax.axhline(3.5, color=GRIS, lw=1, ls="--")          # separador: arriba peso estable, abajo peso depreciado
ax.invert_yaxis()
ax.xaxis.set_major_formatter(mtick.StrMethodFormatter("{x:,.0f} M"))
ax.set_xlabel("Millones MXN al año 40")
ax.set_title("Saldo final al año 40: lo que queda en tu bolsillo tras el ISR de salida", fontsize=11, fontweight="bold")
leyenda_02 = [Patch(facecolor=color, label=etiqueta) for color, etiqueta, _ in ESTILOS]
leyenda_02.append(Patch(facecolor=ROJO, alpha=0.85, label="ISR al retirar"))
ax.set_xlim(0, max(n + i for n, i in zip(netos, imps)) * 1.35)
fig.legend(handles=leyenda_02, loc="lower center", bbox_to_anchor=(0.5, 0.01),
           ncol=3, frameon=False, fontsize=8.5)
fig.tight_layout(rect=(0, 0.11, 1, 1))
fig.savefig(OUT / "02_neto_tras_impuesto.png", dpi=150)
plt.close(fig)

# ---------------------------------------------------- 3. composición del saldo
fig, ax = plt.subplots(figsize=(11.5, 7))
aportes = np.array([ESCENARIOS[i][1]["aportado"] for i in orden]) / M
bonos = np.array([ESCENARIOS[i][1].get("bono", 0.0) for i in orden]) / M
netos = np.array([ESCENARIOS[i][1]["neto"] for i in orden]) / M
crecim = netos - aportes - bonos       # rendimiento neto de cargos, venta e ISR
ax.barh(y, aportes, color=GRIS, label="Aportado (5k/mes; + devolución SAT salvo VUAA mismo bolsillo)")
ax.barh(y, bonos, left=aportes, color=NARANJA, alpha=0.9, label="Bono de Fidelidad (solo PPR)")
ax.barh(y, crecim, left=aportes + bonos, color=VERDE, alpha=0.75, label="Crecimiento neto de cargos, venta e ISR")
for yi, n in zip(y, netos):
    ax.text(n + 0.6, yi, f"neto \\${n:.1f}M", va="center", fontsize=9)
ax.set_yticks(y)
ax.set_yticklabels([etiqueta_barra(i) for i in orden], fontsize=8.5)
ax.axhline(3.5, color=GRIS, lw=1, ls="--")          # separador: arriba peso estable, abajo peso depreciado
ax.invert_yaxis()
ax.xaxis.set_major_formatter(mtick.StrMethodFormatter("{x:,.0f} M"))
ax.set_xlabel("Millones MXN al año 40")
ax.set_title("De dónde sale el neto al retirar (cada barra es dinero que recibes)",
             fontsize=11, fontweight="bold")
ax.legend(loc="upper center", bbox_to_anchor=(0.5, -0.08), ncol=2, frameon=False, fontsize=9)
ax.set_xlim(0, netos.max() * 1.45)
ax.set_ylim(len(orden) - 0.3, -0.7)  # deja la leyenda (abajo, fuera) sin tocar barras
fig.tight_layout()
fig.savefig(OUT / "03_composicion_saldo.png", dpi=150, bbox_inches="tight")
plt.close(fig)

# ---------------------------------------------------- 4. bolsillo vs neto
fig, axes = plt.subplots(1, 2, figsize=(14, 5), sharey=True)
for ax, (titulo, idxs) in zip(axes, grupos):
    x = np.arange(len(idxs))
    bolsillo = [ESCENARIOS[i][1]["bolsillo"] / M for i in idxs]
    neto = [ESCENARIOS[i][1]["neto"] / M for i in idxs]
    ax.bar(x - 0.2, bolsillo, 0.4, color=GRIS, label="Desembolso de tu bolsillo")
    ax.bar(x + 0.2, neto, 0.4, color="#2a9d8f", label="Neto tras impuestos")
    for xi, b, n in zip(x, bolsillo, neto):
        ax.text(xi - 0.2, b + 0.4, f"\\${b:.2f}M", ha="center", fontsize=8.5)
        ax.text(xi + 0.2, n + 0.4, f"\\${n:.1f}M", ha="center", fontsize=9, fontweight="bold")
    ax.set_xticks(x)
    ax.set_xticklabels([estilo[i][1].replace(" que el PPR", "").replace(" · ", "\n") for i in idxs], fontsize=8.5)
    ax.set_title(titulo, fontsize=10)
    ax.yaxis.set_major_formatter(mtick.StrMethodFormatter("{x:,.0f} M"))
axes[0].set_ylabel("Millones MXN")
axes[0].legend(frameon=False, fontsize=9)
fig.suptitle("Lo que pusiste vs lo que te llevas (PPR: el SAT financia 550k; VUAA mismos aportes: todo sale de ti)",
             fontsize=11, fontweight="bold")
fig.tight_layout()
fig.savefig(OUT / "04_bolsillo_vs_neto.png", dpi=150)
plt.close(fig)

# ------------------------------------------------------------------ Excel
wb = Workbook()
encabezado = Font(bold=True, color="FFFFFF")
relleno = PatternFill("solid", fgColor="1F4E79")

ws = wb.active
ws.title = "Supuestos"
supuestos = [
    ("Supuesto", "Valor", "Fuente"),
    ("Aportación mensual", "5,000 MXN", "Usuario"),
    ("Horizonte", "40 años (480 meses), retiro total al final", "Usuario"),
    ("Aportación total", "PPR y VUAA mismos aportes: 2,950,368 MXN (5k×480 + 39×14,112). VUAA mismo bolsillo: 2,400,000 (5k×480)",
     "Dos equivalencias: mismos aportes (la del usuario) y mismo desembolso de bolsillo"),
    ("Devolución SAT anual", "14,112 MXN = 60,000 × 23.52% (llega y se invierte cada junio)", "Tarifa ISR 2026 anual; tope 10% ingreso (Art. 151)"),
    ("Salario bruto", "600,000 MXN/año (asalariado, supuesto)", "Usuario"),
    ("Rendimiento S&P 500", "10% nominal anual USD (histórico con dividendos)", "Supuesto del usuario; real ~6.5-7%"),
    ("Tipo de cambio", "A peso estable (10% en MXN) / B depreciación 1.5% anual (1.10×1.015−1 = 11.65% en MXN)", "Promedio de muy largo plazo USD/MXN"),
    ("Inflación", "4% anual (crece UDI, UMA y referencia del bono)", "Supuesto"),
    ("Cargos PPR (tope)", "1.5% trimestral s/Fondo Inicial (18m) y s/Fondo de Bono + 0.1% mensual s/todo el fondo + 25 UDIS/mes, todo +16% IVA",
     "CG Allianz 3.12 y 3.10.3, CNSF S0003-0247-2018"),
    ("Cargos PPR (típico)", "Mitad de los topes (la carátula puede cobrar menos)", "Supuesto de sensibilidad"),
    ("Bono de Fidelidad", "75% de la aportación comprometida del primer año (45,000 en total), generado pro rata (150/mes); rinde inflación+5% (tope 9%) menos cargos; acreditado año 25. Se supone instrucción para reinvertirlo al rendimiento de mercado; CG indica renta fija corta por defecto",
     "CG Allianz 3.6 y 3.10.3; folleto OptiMaxx plus p.2 y p.5"),
    ("ISR salida PPR", "Caso base: retiro único y toda la exención disponible, hasta 90 UMA elevadas al año (3.85M en 2026, crece con la UMA); tarifa anual sobre el excedente. Otras pensiones pueden consumir la exención; ver sensibilidad fiscal",
     "CG 3.21; regla 3.17.6 RMF; art. 171 RLISR"),
    ("ISR salida VUAA", "10% definitivo sobre la ganancia; costo de adquisición actualizado por inflación", "Art. 129 LISR"),
    ("Tarifa ISR", "2026 congelada nominalmente 40 años (simplificación; indexarla favorece al PPR 0.1-0.7M)", "Anexo 8 RMF 2026"),
    ("VUAA costos", "Corretaje supuesto 0.25% + IVA por operación + OCF 0.07% anual", "Vanguard México (OCF 2026); corretaje supuesto"),
    ("UDI / UMA 2026", "8.826356 MXN / 117.31 MXN diaria", "Banxico / INEGI (sept-2026)"),
]
for fila in supuestos:
    ws.append(fila)
for fila in range(1, len(supuestos) + 1):
    ws.row_dimensions[fila].height = 58 if fila > 1 else 26
    for cell in ws[fila]:
        cell.alignment = Alignment(wrap_text=True, vertical="center")
for c in range(1, 4):
    ws.cell(1, c).font = encabezado
    ws.cell(1, c).fill = relleno
ws.column_dimensions["A"].width = 26
ws.column_dimensions["B"].width = 62
ws.column_dimensions["C"].width = 42

ws = wb.create_sheet("Resultados")
ws.append(["Escenario", "Bolsillo (MXN)", "Aportado (MXN)", "Cargos y costos totales (MXN)", "Bono año 40 (MXN)",
           "Saldo año 40 (MXN)", "ISR salida (MXN)", "Neto tras ISR (MXN)",
           "VUAA corretaje compra (MXN)", "VUAA OCF (MXN)", "VUAA corretaje venta (MXN)",
           "Neto en pesos de 2026 (MXN)"])
for nombre, r in ESCENARIOS:
    ppr = "fees" in r
    com = sum(r["fees"].values()) if ppr else r["costos_total"]
    ws.append([nombre, r["bolsillo"], r["aportado"], com, r.get("bono", 0.0), r["saldo"], r["impuesto"], r["neto"],
               None if ppr else r["corretaje_compra"], None if ppr else r["ocf"],
               None if ppr else r["corretaje_venta"], r["neto_hoy"]])
for c in range(1, 13):
    ws.cell(1, c).font = encabezado
    ws.cell(1, c).fill = relleno
    ws.cell(1, c).alignment = Alignment(wrap_text=True)
ws.row_dimensions[1].height = 48
for c in range(2, 13):
    for fila in range(2, len(ESCENARIOS) + 2):
        ws.cell(fila, c).number_format = "#,##0"
ws.column_dimensions["A"].width = 38
for c in range(2, 13):
    ws.column_dimensions[get_column_letter(c)].width = 19

for (titulo, idxs), hoja in zip(grupos, ["Flujo peso estable", "Flujo peso -1.5% anual"]):
    ws = wb.create_sheet(hoja)
    ws.append(["Año", "Aportado acumulado (5k/mes + devolución)", "Bolsillo acumulado (5k/mes)",
               "PPR tope (sin bono, antes de retiro)", "PPR mitad (sin bono, antes de retiro)", "VUAA mismos aportes", "VUAA mismo bolsillo"])
    for a in range(40):
        ws.append([a + 1, APORT_M * 12 * (a + 1) + REFUND * a, APORT_M * 12 * (a + 1)]
                  + [ESCENARIOS[i][1]["serie"][a] for i in idxs])
    for c in range(1, 8):
        ws.cell(1, c).font = encabezado
        ws.cell(1, c).fill = relleno
        ws.cell(1, c).alignment = Alignment(wrap_text=True)
    ws.row_dimensions[1].height = 50
    for c in range(2, 8):
        for fila in range(2, 42):
            ws.cell(fila, c).number_format = "#,##0"
    ws.column_dimensions["A"].width = 8
    for c in range(2, 8):
        ws.column_dimensions[get_column_letter(c)].width = 20

wb.save(BASE / "outputs" / "resultados_ppr_vs_vuaa.xlsx")
print("OK: gráficas en", OUT, "y Excel en", BASE / "outputs" / "resultados_ppr_vs_vuaa.xlsx")
