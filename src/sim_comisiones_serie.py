"""Gráfica 06: cargos cobrados durante el plazo vs su costo real.

Todo sale de sim_ppr/sim_vuaa (una sola implementación del modelo). Las áreas son
los cargos nominales cobrados al PPR (tope), acumulados año con año. Las líneas son
el costo real: el saldo que se deja de tener por los cargos (el mismo plan con
cargos en cero menos con cargos), que incluye lo que esos pesos habrían rendido.
Escenario 5k/mes + devolución SAT (14,112 cada junio), igual que el resto del ejercicio.
"""
import json
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.ticker as mtick
import numpy as np
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from sim_ppr_vs_vuaa import FX, MESES, sim_ppr, sim_vuaa

BASE = Path(__file__).resolve().parent.parent   # raíz del proyecto
OUT = BASE / "outputs" / "graficas"
plt.rcParams.update({"font.size": 10, "axes.spines.top": False, "axes.spines.right": False})
M = 1e6

if __name__ == "__main__":
    anios = np.arange(1, MESES // 12 + 1)
    datos, resumen = {}, {}
    fig, axes = plt.subplots(1, 2, figsize=(13, 5.6), sharey=True)
    for ax, (g, fx) in zip(axes, FX):
        tope, mitad, sin = sim_ppr(g, 1.0), sim_ppr(g, 0.5), sim_ppr(g, 0.0)
        vu, vu0 = sim_vuaa(g), sim_vuaa(g, costos=0.0)
        c = tope["serie_cargos"]
        cobrado = {"PPR tope": np.sum([c[k] for k in c], axis=0),
                   "PPR mitad": np.sum([mitad["serie_cargos"][k] for k in c], axis=0),
                   "VUAA": np.array(vu["serie_costos"])}
        real = {"PPR tope": np.subtract(sin["serie"], tope["serie"]),
                "PPR mitad": np.subtract(sin["serie"], mitad["serie"]),
                 "VUAA": np.subtract(vu0["serie"], vu["serie"])}
        # Las series anuales son previas a la venta; el extremo final incluye su corretaje.
        cobrado["VUAA"][-1] += vu["corretaje_venta"]
        real["VUAA"][-1] = vu0["saldo"] - vu["saldo"]
        datos[fx] = {**{k: list(v) for k, v in c.items()},
                     **{f"cobrado {k}": list(v) for k, v in cobrado.items()},
                     **{f"costo real {k}": list(v) for k, v in real.items()},
                     "bono_acreditado": tope["bono_acreditado"]}

        ax.stackplot(anios, np.array(c["administrativo"]) / M, np.array(c["gestion"]) / M, np.array(c["fijo"]) / M,
                     labels=["Cobrado tope: administrativo (1.5% trim. s/Fondo Inicial)",
                             "Cobrado tope: gestión (0.1% mensual s/fondo total)", "Cobrado tope: fijo (25 UDIS/mes)"],
                     colors=["#1f77b4", "#ff7f0e", "#9467bd"], alpha=0.8)
        ax.plot(anios, real["PPR tope"] / M, color="#08306b", lw=2.4, label="Costo real PPR tope: saldo que dejas de tener")
        ax.plot(anios, real["PPR mitad"] / M, color="#08306b", lw=1.8, ls="--", label="Costo real PPR mitad de tope")
        ax.plot(anios, real["VUAA"] / M, color="#2ca02c", lw=2.2, label="Costo real VUAA (compra + OCF + venta al año 40)")
        resumen[ax] = "Al año 40, costo real (cobrado):\n" + "\n".join(
            f"{k}: \\${real[k][-1] / M:.1f}M (\\${cobrado[k][-1] / M:.1f}M)" for k in real)
        ax.set_title(fx, fontsize=10.5)
        ax.set_xlabel("Años")
        ax.set_xlim(1, 40)
        ax.yaxis.set_major_formatter(mtick.StrMethodFormatter("{x:,.0f} M"))
        ax.tick_params(labelleft=True)
    axes[0].set_ylabel("MXN nominales acumulados (millones)")
    axes[0].legend(loc="upper left", frameon=False, fontsize=8.5)
    ymax = axes[0].get_ylim()[1]
    axes[1].text(1.5, 0.95 * ymax,       # el bono no depende del tipo de cambio: una sola nota
                 f"Fondo de Bono acreditado al año 25: \\${datos[FX[0][1]]['bono_acreditado'] / 1e3:.0f}k\n"
                 "(bonificación parcial de cargos, CG 3.6)", fontsize=8.5, color="#d62728", va="top")
    for ax, y in zip(axes, (0.65, 0.78)):
        ax.text(1.5, y * ymax, resumen[ax], fontsize=8.5, color="#08306b", va="top")
    fig.suptitle("Cargos cobrados vs costo real (saldo que dejas de tener) — 5,000 MXN/mes + devolución SAT",
                 fontsize=11.5, fontweight="bold")
    fig.tight_layout()
    fig.savefig(OUT / "06_comisiones_durante_plazo.png", dpi=150)
    plt.close(fig)

    (BASE / "outputs" / "resultados_comisiones_serie.json").write_text(
        json.dumps(datos, ensure_ascii=False), encoding="utf-8")

    xlsx = BASE / "outputs" / "resultados_ppr_vs_vuaa.xlsx"
    wb = load_workbook(xlsx)
    if "Comisiones" in wb.sheetnames:
        del wb["Comisiones"]
    ws = wb.create_sheet("Comisiones")
    enc, rell = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E79")
    ws.append(["Año (peso estable)", "PPR tope cobrado: administrativo", "PPR tope cobrado: gestión",
               "PPR tope cobrado: fijo", "PPR tope cobrado: total", "PPR mitad cobrado: total", "VUAA cobrado: compra+OCF+venta al año 40",
               "Costo real PPR tope (saldo perdido)", "Costo real PPR mitad", "Costo real VUAA"])
    d = datos[FX[0][1]]
    for a in range(40):
        ws.append([a + 1, d["administrativo"][a], d["gestion"][a], d["fijo"][a], d["cobrado PPR tope"][a],
                   d["cobrado PPR mitad"][a], d["cobrado VUAA"][a], d["costo real PPR tope"][a],
                   d["costo real PPR mitad"][a], d["costo real VUAA"][a]])
    for c in range(1, 11):
        ws.cell(1, c).font = enc
        ws.cell(1, c).fill = rell
        ws.cell(1, c).alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 48
    for fila in range(2, 42):
        for c in range(2, 11):
            ws.cell(fila, c).number_format = "#,##0"
    ws.column_dimensions["A"].width = 10
    for c in range(2, 11):
        ws.column_dimensions[get_column_letter(c)].width = 19
    wb.save(xlsx)
    print("OK: graficas/06_comisiones_durante_plazo.png + hoja 'Comisiones' en el Excel")
