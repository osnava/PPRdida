"""Escenario adicional: regla del 4% en la fase de retiro (sin venta lump-sum).

Acumulación: 40 años (igual que el resto del ejercicio). Retiro: 30 años
(65 → 95 años), retirando cada año el 4% del saldo inicial actualizado por
inflación (regla clásica de Bengen).

Escenarios de la fase de retiro (ambos supuestos cambiarios):
  - VUAA (mismos aportes y mismo bolsillo): no se vende de golpe; cada retiro
    vende un pedazo del portafolio y paga 10% de ISR solo sobre la fracción de
    ganancia de lo vendido. La base de costo se actualiza por inflación (Art. 129
    LISR) y no crece con el mercado. El impuesto no pagado sigue invertido.
  - PPR (tope y mitad): este contrato exige retirar una sola vez y en su
    totalidad (CG 3.21); el neto del año 40 se reinvierte en VUAA con base igual a
    lo invertido; desde ahí paga el mismo 10% sobre sus ganancias nuevas.

Riqueza total en pesos de hoy (deflactor 4% anual, año por año desde 2026) =
retiros netos + legado del año 70 neto del ISR latente (10% de la ganancia no
realizada), para comparar patrimonios líquidos. Sin OCF ni corretaje en la fase.
"""
import json
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from sim_ppr_vs_vuaa import FX, INFLACION, sim_ppr, sim_vuaa

BASE = Path(__file__).resolve().parent.parent   # raíz del proyecto
ANIOS_RETIRO = 30
TASA_ISR = 0.10


def drawdown(s0: float, b0: float, r_nom: float, infl: float = INFLACION,
             anios: int = ANIOS_RETIRO, pct: float = 0.04):
    """Regla del 4%: retiro anual = pct del saldo inicial, ajustado por inflación.

    Cada retiro paga ISR sobre su fracción de ganancia; la base se actualiza por
    inflación (Art. 129). Devuelve agregados nominales y en pesos de hoy, y la serie.
    """
    s, b = s0, b0
    w1 = pct * s0
    tot = dict.fromkeys(["total_retirado", "total_isr", "neto_hoy", "isr_hoy"], 0.0)
    serie = []
    for t in range(anios):
        if s <= 0:
            break
        w = min(w1 * (1 + infl) ** t, s)
        isr = TASA_ISR * max(0.0, 1 - b / s) * w
        b -= b / s * w
        s = (s - w) * (1 + r_nom)
        b *= 1 + infl
        d = (1 + infl) ** (40 + t)        # el retiro t+1 ocurre 40 + t años después de 2026
        tot["total_retirado"] += w
        tot["total_isr"] += isr
        tot["neto_hoy"] += (w - isr) / d
        tot["isr_hoy"] += isr / d
        serie.append({"anio": 41 + t, "retiro": w, "isr": isr, "neto": w - isr, "saldo": s})
    d70 = (1 + infl) ** (40 + anios)
    legado_hoy = max(0.0, s) / d70
    latente_hoy = TASA_ISR * max(0.0, s - b) / d70
    retiro_anio1_neto = serie[0]["neto"] if serie else 0.0
    return {
        "retiro_anio1": w1, "retiro_anio1_hoy": w1 / (1 + infl) ** 40, **tot,
        "retiro_anio1_neto": retiro_anio1_neto,
        "retiro_anio1_neto_hoy": retiro_anio1_neto / (1 + infl) ** 40,
        "total_neto": tot["total_retirado"] - tot["total_isr"], "saldo_final": max(0.0, s),
        "legado_hoy": legado_hoy, "latente_hoy": latente_hoy,
        "riqueza_hoy": tot["neto_hoy"] + legado_hoy - latente_hoy, "serie": serie,
    }


def filas_retiro():
    filas = []
    for g, fx in FX:
        casos = []
        for nombre, con_refund in (("VUAA mismos aportes", True), ("VUAA mismo bolsillo", False)):
            v = sim_vuaa(g, con_refund)
            casos.append((nombre, v["saldo_pre_venta"], v["base"]))
        for nombre, fs in (("PPR tope → reinvertido", 1.0), ("PPR mitad → reinvertido", 0.5)):
            neto = sim_ppr(g, fs)["neto"]
            casos.append((nombre, neto, neto))
        for nombre, s0, b0 in casos:
            d = drawdown(s0, b0, g)
            d["escenario"] = f"{nombre} · {fx}"
            filas.append(d)
    return filas


if __name__ == "__main__":
    if sys.stdout.encoding and sys.stdout.encoding.lower() != "utf-8":
        sys.stdout.reconfigure(encoding="utf-8")
    filas = filas_retiro()
    M = 1e6
    print(f"{'escenario':<42}{'retiro neto año 1 (hoy)':>26}{'ISR fase (hoy)':>16}{'legado neto (hoy)':>19}{'riqueza total (hoy)':>21}")
    for d in filas:
        print(f"{d['escenario']:<42}{d['retiro_anio1_neto_hoy']/M:>25.3f}M{d['isr_hoy']/M:>15.2f}M"
              f"{(d['legado_hoy'] - d['latente_hoy'])/M:>18.2f}M{d['riqueza_hoy']/M:>20.2f}M")

    # hoja nueva en el Excel existente
    xlsx = BASE / "outputs" / "resultados_ppr_vs_vuaa.xlsx"
    wb = load_workbook(xlsx)
    if "Retiro 4%" in wb.sheetnames:
        del wb["Retiro 4%"]
    ws = wb.create_sheet("Retiro 4%")
    enc = Font(bold=True, color="FFFFFF")
    rell = PatternFill("solid", fgColor="1F4E79")
    ws.append(["Escenario (fase retiro 30 años, regla del 4%)",
               "Retiro bruto año 1 nominal (MXN)", "Retiro bruto año 1 en pesos de hoy (MXN)",
               "ISR fase retiro nominal (MXN)", "ISR fase retiro en pesos de hoy (MXN)",
               "Retiros netos nominal (MXN)", "Retiros netos en pesos de hoy (MXN)",
               "Saldo año 70 nominal (MXN)", "Legado año 70 en pesos de hoy (MXN)",
               "ISR latente del legado en pesos de hoy (MXN)",
               "Riqueza total en pesos de hoy (MXN)",
               "Retiro neto año 1 nominal (MXN)", "Retiro neto año 1 en pesos de hoy (MXN)"])
    for d in filas:
        ws.append([d["escenario"], d["retiro_anio1"], d["retiro_anio1_hoy"], d["total_isr"], d["isr_hoy"],
                   d["total_neto"], d["neto_hoy"], d["saldo_final"], d["legado_hoy"], d["latente_hoy"],
                   d["riqueza_hoy"], d["retiro_anio1_neto"], d["retiro_anio1_neto_hoy"]])
    for c in range(1, 14):
        ws.cell(1, c).font = enc
        ws.cell(1, c).fill = rell
        ws.cell(1, c).alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 55
    for fila in range(2, len(filas) + 2):
        for c in range(2, 14):
            ws.cell(fila, c).number_format = "#,##0"
    ws.column_dimensions["A"].width = 42
    for c in range(2, 14):
        ws.column_dimensions[get_column_letter(c)].width = 19
    wb.save(xlsx)

    (BASE / "outputs" / "resultados_retiro_4pct.json").write_text(
        json.dumps(filas, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print("\nOK: hoja 'Retiro 4%' añadida al Excel y resultados_retiro_4pct.json")
