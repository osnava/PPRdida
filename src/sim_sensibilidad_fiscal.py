"""Sensibilidad del retiro único del PPR; no sustituye el cálculo fiscal personal.

Acumulación: 5,000 MXN/mes durante 40 años, ingreso 600,000/año, devolución
calculada con la tarifa ISR 2026 congelada nominalmente. Rendimientos: 10% MXN
con peso estable y 11.65% MXN con depreciación compuesta del 1.5%; inflación
4%; cargos Allianz al tope o a mitad. Compara 0%, 50% y 100% de la exención
de 90 UMA anuales consumida por otras pensiones. La opción Art. 95 usa otros
ingresos gravables y último sueldo mensual explícitos; es una sensibilidad,
pendiente de validación fiscal del caso real. Marco: VUAA con mismo bolsillo.
"""
import argparse
import json
from dataclasses import replace
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from sim_ppr_vs_vuaa import EXENCION_UMAS, FX, Parametros, sim_ppr, sim_vuaa

BASE = Path(__file__).resolve().parent.parent


def filas_fiscales(otros_ingresos_gravables: float = 600_000.0,
                  ultimo_sueldo_mensual: float = 50_000.0) -> dict:
    """Devuelve escenarios de exención disponible y dos métodos de ISR."""
    base = Parametros()
    exencion_total = EXENCION_UMAS * base.uma_diaria * (1 + base.inflacion) ** (base.meses / 12)
    filas = []
    for tasa_mercado, fx in FX:
        referencia = sim_vuaa(tasa_mercado, con_refund=False, config=base)
        for escala, cargo in ((1.0, "tope"), (0.5, "mitad")):
            for fraccion in (0.0, 0.5, 1.0):
                for metodo in ("tarifa", "art95"):
                    config = replace(
                        base,
                        exencion_ya_usada_mxn=exencion_total * fraccion,
                        metodo_isr_retiro=metodo,
                        otros_ingresos_gravables=otros_ingresos_gravables if metodo == "art95" else 0.0,
                        ultimo_sueldo_mensual=ultimo_sueldo_mensual if metodo == "art95" else 0.0,
                    )
                    r = sim_ppr(tasa_mercado, escala, config=config)
                    filas.append({
                        "fx": fx, "rendimiento_anual_mxn": tasa_mercado,
                        "cargos": cargo, "fee_scale": escala,
                        "fraccion_exencion_consumida": fraccion,
                        "metodo_isr": metodo, "exencion_disponible": r["exencion_disponible"],
                        "saldo": r["saldo"], "isr_ppr": r["impuesto"],
                        "neto_ppr": r["neto"], "neto_ppr_hoy": r["neto_hoy"],
                        "neto_vuaa_mismo_bolsillo": referencia["neto"],
                        "diferencia_ppr_menos_vuaa": r["neto"] - referencia["neto"],
                    })
    return {
        "supuestos": {
            "anio_tarifa": 2026, "tarifa_congelada_nominalmente": True,
            "inflacion": base.inflacion, "aportacion_mensual": base.aportacion_mensual,
            "ingreso_anual_aportacion": base.ingreso_anual,
            "exencion_total_nominal_al_retiro": exencion_total,
            "otros_ingresos_gravables_art95": otros_ingresos_gravables,
            "ultimo_sueldo_mensual_art95": ultimo_sueldo_mensual,
            "marco_comparacion": "mismo desembolso de bolsillo",
            "alcance": "sensibilidad aislada: no calcula el ISR conjunto de otras pensiones",
        },
        "resultados": filas,
    }


def agregar_hoja(datos: dict) -> None:
    """Guarda los 24 cruces en el libro generado por el pipeline."""
    path = BASE / "outputs" / "resultados_ppr_vs_vuaa.xlsx"
    wb = load_workbook(path)
    if "Sensibilidad fiscal" in wb.sheetnames:
        del wb["Sensibilidad fiscal"]
    ws = wb.create_sheet("Sensibilidad fiscal")
    ws.append([
        "Tipo de cambio", "Rendimiento MXN", "Cargos PPR", "Exención usada por otras pensiones",
        "Método ISR", "Exención disponible (MXN)", "Saldo PPR (MXN)", "ISR PPR (MXN)",
        "Neto PPR (MXN)", "Neto PPR pesos 2026 (MXN)", "VUAA mismo bolsillo (MXN)",
        "Diferencia PPR − VUAA (MXN)",
    ])
    for r in datos["resultados"]:
        ws.append([
            r["fx"], r["rendimiento_anual_mxn"], r["cargos"], r["fraccion_exencion_consumida"],
            r["metodo_isr"], r["exencion_disponible"], r["saldo"], r["isr_ppr"],
            r["neto_ppr"], r["neto_ppr_hoy"], r["neto_vuaa_mismo_bolsillo"],
            r["diferencia_ppr_menos_vuaa"],
        ])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 42
    for row in ws.iter_rows(min_row=2):
        for col in (2, 4):
            row[col - 1].number_format = "0.0%"
        for cell in row[5:]:
            cell.number_format = "#,##0"
    ws.column_dimensions["A"].width = 25
    for col in range(2, 13):
        ws.column_dimensions[get_column_letter(col)].width = 21
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:L{ws.max_row}"
    ws.append(["Art. 95: otros ingresos gravables y último sueldo son hipótesis configurables; no incluye el ISR conjunto de otras pensiones."])
    ws.append([f"Hipótesis Art. 95: otros ingresos {datos['supuestos']['otros_ingresos_gravables_art95']:,.0f} MXN; último sueldo {datos['supuestos']['ultimo_sueldo_mensual_art95']:,.0f} MXN/mes. Tarifa 2026 congelada."])
    for row_no in (ws.max_row - 1, ws.max_row):
        ws.merge_cells(start_row=row_no, start_column=1, end_row=row_no, end_column=12)
        ws.cell(row_no, 1).alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[row_no].height = 28
    wb.save(path)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--otros-ingresos", type=float, default=600_000.0,
                        help="Ingresos gravables adicionales en el año del retiro para Art. 95")
    parser.add_argument("--ultimo-sueldo", type=float, default=50_000.0,
                        help="Último sueldo mensual ordinario para Art. 95")
    args = parser.parse_args()
    if args.otros_ingresos < 0 or args.ultimo_sueldo <= 0:
        parser.error("Otros ingresos debe ser >= 0 y último sueldo > 0")
    datos = filas_fiscales(args.otros_ingresos, args.ultimo_sueldo)
    path = BASE / "outputs" / "resultados_sensibilidad_fiscal.json"
    path.write_text(json.dumps(datos, ensure_ascii=False, indent=2), encoding="utf-8")
    agregar_hoja(datos)
    print(f"OK: {len(datos['resultados'])} sensibilidades -> {path}")
    for r in datos["resultados"]:
        if r["metodo_isr"] == "tarifa":
            print(f"{r['fx']:<21} PPR {r['cargos']:<5} exención usada {r['fraccion_exencion_consumida']:.0%}: "
                  f"neto {r['neto_ppr']/1e6:.2f}M vs VUAA mismo bolsillo "
                  f"{r['neto_vuaa_mismo_bolsillo']/1e6:.2f}M")
