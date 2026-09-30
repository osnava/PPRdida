"""Regenera el análisis completo en el orden requerido y verifica sus artefactos.

Uso: .venv/Scripts/python.exe src/generar_todo.py
Los supuestos fiscales de los resultados base siguen congelados en 2026.
"""
import json
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook

from sim_ppr_vs_vuaa import INFLACION, MESES

BASE = Path(__file__).resolve().parent.parent
SCRIPTS = (
    "sim_ppr_vs_vuaa.py",
    "sim_tarifa_folleto.py",
    "grafica_tarifa_folleto.py",
    "graficas_tablas.py",
    "sim_retiro_4pct.py",
    "grafica_retiro.py",
    "sim_comisiones_serie.py",
    "sim_comisiones.py",
    "sim_sensibilidad_fiscal.py",
    "auditoria.py",
)
HOJAS = {
    "Supuestos", "Resultados", "Flujo peso estable", "Flujo peso -1.5% anual",
    "Retiro 4%", "Comisiones", "Sensibilidad fiscal",
}


def verificar():
    salida = BASE / "outputs"
    libro = load_workbook(salida / "resultados_ppr_vs_vuaa.xlsx", read_only=True, data_only=True)
    try:
        if set(libro.sheetnames) != HOJAS:
            raise RuntimeError(f"Hojas del Excel incompletas: {libro.sheetnames}")
        ws = libro["Resultados"]
        for fila in ws.iter_rows(min_row=2, values_only=True):
            if fila[0].startswith("VUAA"):
                if abs(fila[3] - sum(fila[8:11])) > 0.01:
                    raise RuntimeError(f"Los costos VUAA no concilian en {fila[0]}")
            if abs(fila[7] - fila[11] * ((1 + INFLACION) ** (MESES / 12))) > 0.01:
                raise RuntimeError(f"El neto real no concilia en {fila[0]}")
        ws_fiscal = libro["Sensibilidad fiscal"]
        if ws_fiscal.max_row != 27:
            raise RuntimeError(f"La hoja fiscal no tiene 24 cruces: {ws_fiscal.max_row} filas")
    finally:
        libro.close()
    resultados = json.loads((salida / "resultados_ppr_vs_vuaa.json").read_text(encoding="utf-8"))
    if len(resultados) != 8:
        raise RuntimeError("Faltan escenarios en el JSON principal")
    folleto = json.loads((salida / "resultados_tarifa_folleto.json").read_text(encoding="utf-8"))
    if len(folleto["escenarios"]) != 2:
        raise RuntimeError("Faltan escenarios de la tarifa del folleto")
    for fx, caso in folleto["escenarios"].items():
        if not caso["ppr_tope_neto"] < caso["ppr_folleto_neto"] < caso["ppr_mitad_neto"]:
            raise RuntimeError(f"La sensibilidad del folleto no concilia en {fx}")
    sensibilidad = json.loads((salida / "resultados_sensibilidad_fiscal.json").read_text(encoding="utf-8"))
    if len(sensibilidad["resultados"]) != 24:
        raise RuntimeError("Faltan escenarios fiscales")
    for i in range(1, 8):
        if not list((salida / "graficas").glob(f"{i:02d}_*.png")):
            raise RuntimeError(f"Falta la gráfica {i:02d}")
    print("VERIFICADO: 7 hojas, 8 escenarios base, 2 escenarios de folleto, 24 sensibilidades y 7 gráficas")


if __name__ == "__main__":
    if sys.stdout.encoding and sys.stdout.encoding.lower() != "utf-8":
        sys.stdout.reconfigure(encoding="utf-8")
    for script in SCRIPTS:
        print(f"\n==> {script}", flush=True)
        subprocess.run([sys.executable, "-B", str(BASE / "src" / script)], cwd=BASE, check=True)
    verificar()
